import os

from docx import Document

from fpdf import FPDF


# -----------------------------------
# Output Folder
# -----------------------------------

OUTPUT_FOLDER = "generated_documents"


os.makedirs(
    OUTPUT_FOLDER,
    exist_ok=True
)


# -----------------------------------
# TXT
# -----------------------------------

def create_txt(document_text):

    file_path = os.path.join(
        OUTPUT_FOLDER,
        "LegalEase_Document.txt"
    )


    with open(
        file_path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(document_text)


    return file_path


# -----------------------------------
# DOCX
# -----------------------------------

def create_docx(document_text):

    file_path = os.path.join(
        OUTPUT_FOLDER,
        "LegalEase_Document.docx"
    )


    document = Document()


    # Split document into lines
    lines = document_text.split("\n")


    for line in lines:

        line = line.strip()


        if not line:
            continue


        # First line as title
        if len(document.paragraphs) == 0:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True

        else:

            document.add_paragraph(line)


    document.save(file_path)


    return file_path


# -----------------------------------
# PDF
# -----------------------------------

def create_pdf(document_text):

    file_path = os.path.join(
        OUTPUT_FOLDER,
        "LegalEase_Document.pdf"
    )


    pdf = FPDF()


    pdf.set_auto_page_break(
        auto=True,
        margin=15
    )


    pdf.add_page()


    pdf.set_font(
        "Arial",
        size=12
    )


    lines = document_text.split("\n")


    for line in lines:

        line = line.strip()


        if not line:

            pdf.ln(5)

            continue


        pdf.multi_cell(
            0,
            8,
            line
        )


    pdf.output(
        file_path
    )


    return file_path
